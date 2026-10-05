"""Ingestão do material do projeto (Camada Leitora, parte 1).

Fluxo principal: PDF com camada de texto, exibido no pdf.js.
Plano B (formato só é conhecido no Dia 1):
  - DOCX            -> conversão local para PDF com LibreOffice (soffice).
  - PDF escaneado   -> OCR local com Tesseract; exibido como texto por página.
  - TXT / MD        -> parágrafos numerados.
  - CSV / XLSX      -> cada linha vira um parágrafo numerado.
Nada sai da máquina: nenhuma conversão usa serviço externo.
"""

from __future__ import annotations

import csv
import io
import logging
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path

import pymupdf

from app.domain.clock import iso, now_utc
from app.domain.errors import InvalidInput, UnsupportedDocument
from app.domain.hashing import sha256_bytes
from app.domain.models import DocumentContent, DocumentInfo, Paragraph
from app.reader.text import collapse_ws

log = logging.getLogger("ingestion")

PDF_NAME = "documento.pdf"
ORIGINAL_PREFIX = "original"
MIN_CHARS_PER_PAGE = 25
MAX_PAGES = 500
EXTENSIONS = {".pdf": "pdf", ".docx": "docx", ".txt": "txt", ".md": "md", ".csv": "csv", ".xlsx": "xlsx"}


@dataclass(frozen=True)
class IngestedDocument:
    info: DocumentInfo
    content: DocumentContent


def ingest(raw: bytes, filename: str, dest: Path, max_bytes: int) -> IngestedDocument:
    if not raw:
        raise InvalidInput("O arquivo enviado está vazio.")
    if len(raw) > max_bytes:
        raise InvalidInput(f"O arquivo excede o limite de {max_bytes // (1024 * 1024)} MB.")
    suffix = Path(filename).suffix.lower()
    fmt = EXTENSIONS.get(suffix)
    if fmt is None:
        raise UnsupportedDocument(f"Formato '{suffix or 'sem extensão'}' não suportado. Envie PDF, DOCX, TXT, MD, CSV ou XLSX.")

    dest.mkdir(parents=True, exist_ok=True)
    for old in dest.iterdir():
        if old.is_file():
            old.unlink()
    (dest / f"{ORIGINAL_PREFIX}{suffix}").write_bytes(raw)

    warnings: list[str] = []
    ocr = False
    if fmt == "pdf":
        content, view, ocr = _from_pdf(raw, dest, warnings)
    elif fmt == "docx":
        pdf_bytes = _docx_to_pdf(raw)
        content, view, ocr = _from_pdf(pdf_bytes, dest, warnings)
        warnings.append("Documento convertido de DOCX para PDF localmente; a paginação segue a conversão.")
    elif fmt in ("txt", "md"):
        content, view = _from_text(_decode(raw, warnings)), "text"
    elif fmt == "csv":
        content, view = _from_csv(_decode(raw, warnings)), "text"
    else:
        content, view = _from_xlsx(raw), "text"

    if view == "text" and not content.paragraphs:
        raise InvalidInput("Não foi possível extrair nenhum texto do arquivo.")

    info = DocumentInfo(
        original_filename=Path(filename).name[:200],
        source_format=fmt,  # type: ignore[arg-type]
        view=view,  # type: ignore[arg-type]
        page_count=len(content.pages),
        paragraph_count=len(content.paragraphs),
        sha256=sha256_bytes(raw),
        size_bytes=len(raw),
        uploaded_at=iso(now_utc()),
        ocr_applied=ocr,
        warnings=warnings,
    )
    log.info("document.ingested", extra={"ctx": {"format": fmt, "view": view, "pages": info.page_count, "ocr": ocr}})
    return IngestedDocument(info, content)


# ----------------------------------------------------------------------- PDF
def _from_pdf(raw: bytes, dest: Path, warnings: list[str]) -> tuple[DocumentContent, str, bool]:
    try:
        doc = pymupdf.open(stream=raw, filetype="pdf")
    except Exception as exc:  # noqa: BLE001 - PyMuPDF lança tipos variados
        raise InvalidInput("O PDF está corrompido ou não pôde ser aberto.") from exc
    with doc:
        if doc.needs_pass:
            raise InvalidInput("O PDF é protegido por senha. Envie uma versão sem proteção.")
        if doc.page_count == 0:
            raise InvalidInput("O PDF não tem páginas.")
        if doc.page_count > MAX_PAGES:
            raise InvalidInput(f"O PDF tem {doc.page_count} páginas; o limite é {MAX_PAGES}.")
        pages = [page.get_text("text") for page in doc]
        (dest / PDF_NAME).write_bytes(raw)

        text_chars = sum(len(p.strip()) for p in pages)
        if text_chars >= MIN_CHARS_PER_PAGE * doc.page_count * 0.5:
            empty = [i + 1 for i, p in enumerate(pages) if len(p.strip()) < MIN_CHARS_PER_PAGE]
            if empty:
                warnings.append(f"Páginas com pouco ou nenhum texto selecionável: {_ranges(empty)}.")
            return DocumentContent(pages=pages), "pdf", False

        ocr_pages = _ocr(doc)
    paragraphs = _paragraphs_from_pages(ocr_pages)
    warnings.append(
        "O PDF não tem camada de texto (provavelmente escaneado). O texto foi obtido por OCR local e exibido por página; confira as evidências com atenção."
    )
    return DocumentContent(pages=ocr_pages, paragraphs=paragraphs), "text", True


def _ocr(doc: pymupdf.Document) -> list[str]:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise UnsupportedDocument("PDF escaneado detectado, mas o OCR local não está instalado neste ambiente.") from exc
    if shutil.which("tesseract") is None:
        raise UnsupportedDocument("PDF escaneado detectado, mas o Tesseract não está instalado neste ambiente.")
    langs = pytesseract.get_languages(config="")
    lang = "por" if "por" in langs else "eng"
    pages: list[str] = []
    for page in doc:
        pix = page.get_pixmap(dpi=200)
        image = Image.open(io.BytesIO(pix.tobytes("png")))
        pages.append(pytesseract.image_to_string(image, lang=lang))
    return pages


def _paragraphs_from_pages(pages: list[str]) -> list[Paragraph]:
    out: list[Paragraph] = []
    for page_no, page in enumerate(pages, start=1):
        for block in _split_blocks(page):
            out.append(Paragraph(n=len(out) + 1, text=block, page=page_no))
    return out


# ---------------------------------------------------------------------- DOCX
def _docx_to_pdf(raw: bytes) -> bytes:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        raise UnsupportedDocument("A conversão de DOCX não está disponível neste ambiente (LibreOffice ausente).")
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "entrada.docx"
        src.write_bytes(raw)
        try:
            subprocess.run(
                [soffice, "--headless", "--norestore", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                check=True,
                capture_output=True,
                timeout=120,
                env={"HOME": tmp, "PATH": "/usr/bin:/bin"},
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
            raise InvalidInput("Não foi possível converter o DOCX para PDF. Tente exportar o arquivo como PDF.") from exc
        out = Path(tmp) / "entrada.pdf"
        if not out.exists():
            raise InvalidInput("A conversão do DOCX não gerou um PDF.")
        return out.read_bytes()


# ---------------------------------------------------------------- texto e planilha
def _decode(raw: bytes, warnings: list[str]) -> str:
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    warnings.append("O arquivo não está em UTF-8; foi lido como Latin-1.")
    return raw.decode("latin-1")


def _split_blocks(text: str) -> list[str]:
    blocks = [collapse_ws(b) for b in text.replace("\r\n", "\n").split("\n\n")]
    return [b for b in blocks if b]


def _from_text(text: str) -> DocumentContent:
    paragraphs = [Paragraph(n=i, text=b) for i, b in enumerate(_split_blocks(text), start=1)]
    return DocumentContent(paragraphs=paragraphs)


def _rows_to_paragraphs(rows: list[list[str]]) -> list[Paragraph]:
    rows = [[collapse_ws(str(c)) for c in r] for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return []
    header, body = rows[0], rows[1:] or rows
    out = []
    for row in body:
        if body is rows:
            text = " | ".join(c for c in row if c)
        else:
            text = "; ".join(f"{h or f'coluna {j + 1}'}: {c}" for j, (h, c) in enumerate(zip(header, row, strict=False)) if c)
        out.append(Paragraph(n=len(out) + 1, text=text))
    return out


def _from_csv(text: str) -> DocumentContent:
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return DocumentContent(paragraphs=_rows_to_paragraphs(list(csv.reader(io.StringIO(text), dialect))))


def _from_xlsx(raw: bytes) -> DocumentContent:
    try:
        import openpyxl
    except ImportError as exc:
        raise UnsupportedDocument("Leitura de XLSX indisponível neste ambiente.") from exc
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise InvalidInput("A planilha está corrompida ou não pôde ser aberta.") from exc
    paragraphs: list[Paragraph] = []
    for ws in wb.worksheets:
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        for p in _rows_to_paragraphs(rows):
            paragraphs.append(Paragraph(n=len(paragraphs) + 1, text=f"[{ws.title}] {p.text}"))
    return DocumentContent(paragraphs=paragraphs)


def _ranges(nums: list[int]) -> str:
    parts, start, prev = [], nums[0], nums[0]
    for n in nums[1:] + [None]:  # type: ignore[list-item]
        if n is not None and n == prev + 1:
            prev = n
            continue
        parts.append(str(start) if start == prev else f"{start}–{prev}")
        if n is not None:
            start = prev = n
    return ", ".join(parts)
